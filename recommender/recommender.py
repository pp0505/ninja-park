"""
课程推荐模块 - 读取 Excel 评分表并计算最佳匹配
用法：
    from recommender import ClassRecommender
    rec = ClassRecommender("data/gym.xlsx")
    results = rec.recommend(["协调", "放松"])
"""

from pathlib import Path
from openpyxl import load_workbook

import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config import EXCEL_PATH


class ClassRecommender:
    """读取老板的 Excel 评分表，根据用户选择推荐最佳课程"""

    # 配色和图标池（自动分配给班型）
    COLORS = ["#e63946", "#2a9d8f", "#e9c46a", "#457b9d", "#f4a261", "#6a4c93", "#06d6a0", "#ef476f"]
    EMOJIS = ["🥷", "💪", "🤸", "⚡", "🔥", "🌟", "🎯", "🏆"]

    def __init__(self, excel_path: str | Path = EXCEL_PATH):
        self.excel_path = Path(excel_path)
        self.dimensions = []      # ["上肢力量", "协调", ...]
        self.classes = []          # [{"name": "忍者班", "scores": [7,9,...], "color": ..., "emoji": ...}, ...]
        self._load()

    # ── 读取 Excel ──
    def _load(self):
        if not self.excel_path.exists():
            print(f"⚠️  评分表不存在: {self.excel_path}")
            return

        wb = load_workbook(self.excel_path, data_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            print("⚠️  Excel 数据不足")
            return

        # 第一行：班型名称（跳过第一列）
        header = rows[0]
        class_names = [str(c).strip() for c in header[1:] if c is not None]

        # 后续行：维度名称 + 分数
        self.dimensions = []
        score_matrix = [[] for _ in class_names]

        for row in rows[1:]:
            if row[0] is None:
                continue
            self.dimensions.append(str(row[0]).strip())
            for ci, name in enumerate(class_names):
                val = row[ci + 1] if ci + 1 < len(row) else 0
                score_matrix[ci].append(float(val) if val is not None else 0)

        self.classes = [
            {
                "name": name,
                "scores": score_matrix[i],
                "color": self.COLORS[i % len(self.COLORS)],
                "emoji": self.EMOJIS[i % len(self.EMOJIS)],
            }
            for i, name in enumerate(class_names)
        ]

        print(f"📊 已载入评分表: {len(self.classes)} 个班型, {len(self.dimensions)} 个维度")

    # ── 重新加载（老板换了新 Excel 后调用） ──
    def reload(self, new_path: str | Path | None = None):
        if new_path:
            self.excel_path = Path(new_path)
        self._load()

    # ── 计算推荐 ──
    def recommend(self, selected_dimensions: list[str]) -> list[dict]:
        """
        参数: selected_dimensions - 用户选择的维度名称列表，如 ["协调", "放松"]
        返回: 按匹配度降序排列的班型列表
        """
        # 找到选中维度的索引
        dim_indices = [
            i for i, d in enumerate(self.dimensions)
            if d in selected_dimensions
        ]

        if not dim_indices:
            return []

        max_possible = len(dim_indices) * 10  # 每个维度满分 10

        results = []
        for cls in self.classes:
            total = sum(cls["scores"][i] for i in dim_indices)
            pct = round((total / max_possible) * 100) if max_possible > 0 else 0

            # 各维度的具体得分
            breakdown = {
                self.dimensions[i]: cls["scores"][i]
                for i in dim_indices
            }

            results.append({
                "name": cls["name"],
                "emoji": cls["emoji"],
                "color": cls["color"],
                "total_score": total,
                "match_percentage": pct,
                "breakdown": breakdown,
            })

        results.sort(key=lambda x: x["total_score"], reverse=True)
        return results

    # ── 导出为 API 可用的 JSON 格式 ──
    def to_dict(self) -> dict:
        return {
            "dimensions": self.dimensions,
            "classes": [
                {
                    "name": c["name"],
                    "scores": c["scores"],
                    "color": c["color"],
                    "emoji": c["emoji"],
                }
                for c in self.classes
            ],
        }


# ── 直接运行测试 ──
if __name__ == "__main__":
    rec = ClassRecommender()
    print(f"\n维度: {rec.dimensions}")
    print(f"班型: {[c['name'] for c in rec.classes]}")

    # 模拟用户选择 "协调" 和 "放松"
    results = rec.recommend(["协调", "放松"])
    print("\n🎯 推荐结果（选择: 协调 + 放松）:")
    for r in results:
        print(f"  {r['emoji']} {r['name']}: {r['match_percentage']}% (得分 {r['total_score']})")
        for dim, score in r["breakdown"].items():
            print(f"      {dim}: {score}/10")
